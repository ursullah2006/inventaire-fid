import os
import io
import sqlite3
import smtplib
from email.message import EmailMessage
from datetime import datetime, timedelta, timezone, date
from functools import wraps

from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "cle-secrete-locale-a-changer")

DATABASE_URL = os.environ.get("DATABASE_URL")
USE_PG = bool(DATABASE_URL)
DB_NAME = "inventaire.db"

if USE_PG:
    import psycopg2
    import psycopg2.extras
    DBIntegrityError = psycopg2.IntegrityError
else:
    DBIntegrityError = sqlite3.IntegrityError

STATUTS = ["En service", "En réparation", "Affecté", "Hors service", "Perdu/volé"]
FUSEAU_MADAGASCAR = timezone(timedelta(hours=3))

# Configuration SMTP pour l'envoi d'emails (à remplir via les variables d'environnement)
SMTP_HOST = os.environ.get("SMTP_HOST", "")
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587"))
SMTP_USER = os.environ.get("SMTP_USER", "")
SMTP_PASS = os.environ.get("SMTP_PASS", "")
EMAIL_FROM = os.environ.get("EMAIL_FROM", SMTP_USER)
SMTP_ACTIF = bool(SMTP_HOST and SMTP_USER and SMTP_PASS)


class PGConnection:
    """Enveloppe PostgreSQL qui se comporte comme la connexion SQLite."""

    def __init__(self, conn):
        self.conn = conn

    def execute(self, sql, params=()):
        cur = self.conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        try:
            cur.execute(sql.replace("?", "%s"), params)
        except Exception:
            self.conn.rollback()
            raise
        return cur

    def commit(self):
        self.conn.commit()

    def close(self):
        self.conn.close()


def get_db_connection():
    if USE_PG:
        return PGConnection(psycopg2.connect(DATABASE_URL, connect_timeout=15))
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def inserer(conn, sql, params):
    """Insère une ligne et renvoie son id (SQLite et PostgreSQL)."""
    if USE_PG:
        cur = conn.execute(sql + " RETURNING id", params)
        return cur.fetchone()["id"]
    return conn.execute(sql, params).lastrowid


def init_db():
    cle = "SERIAL PRIMARY KEY" if USE_PG else "INTEGER PRIMARY KEY AUTOINCREMENT"
    reel = "DOUBLE PRECISION" if USE_PG else "REAL"

    conn = get_db_connection()
    conn.execute(f"""
        CREATE TABLE IF NOT EXISTS users (
            id {cle},
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'user'
        )
    """)
    conn.execute(f"""
        CREATE TABLE IF NOT EXISTS categories (
            id {cle},
            nom TEXT NOT NULL UNIQUE,
            icone TEXT DEFAULT '📦'
        )
    """)
    conn.execute(f"""
        CREATE TABLE IF NOT EXISTS materiel (
            id {cle},
            categorie_id INTEGER NOT NULL,
            type TEXT,
            code_immo TEXT NOT NULL UNIQUE,
            prix {reel},
            projet TEXT,
            nom TEXT NOT NULL,
            date_materiel TEXT,
            statut TEXT DEFAULT 'En service',
            responsable TEXT,
            fournisseur TEXT,
            FOREIGN KEY (categorie_id) REFERENCES categories (id) ON DELETE CASCADE
        )
    """)
    conn.execute(f"""
        CREATE TABLE IF NOT EXISTS historique (
            id {cle},
            materiel_id INTEGER,
            code_immo TEXT,
            nom_materiel TEXT,
            action TEXT NOT NULL,
            details TEXT,
            utilisateur TEXT,
            date_action TEXT
        )
    """)
    conn.execute(f"""
        CREATE TABLE IF NOT EXISTS prets (
            id {cle},
            materiel_id INTEGER,
            code_immo TEXT,
            nom_materiel TEXT,
            emprunteur TEXT NOT NULL,
            email_emprunteur TEXT,
            date_debut TEXT,
            date_fin TEXT,
            date_retour TEXT,
            statut TEXT DEFAULT 'En cours'
        )
    """)

    cols = [r["name"] for r in conn.execute("PRAGMA table_info(materiel)").fetchall()]
    if "quantite" not in cols:
        conn.execute("ALTER TABLE materiel ADD COLUMN quantite INTEGER DEFAULT 1")
    if "quantite_initiale" not in cols:
        conn.execute("ALTER TABLE materiel ADD COLUMN quantite_initiale INTEGER DEFAULT 1")

    nb_users = conn.execute("SELECT COUNT(*) AS total FROM users").fetchone()["total"]
    if nb_users == 0:
        conn.execute(
            "INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
            ("admin", generate_password_hash("admin123"), "admin")
        )

    conn.commit()
    conn.close()


def log_historique(materiel_id, code_immo, nom_materiel, action, details=""):
    maintenant = datetime.now(FUSEAU_MADAGASCAR).strftime("%Y-%m-%d %H:%M")
    conn = get_db_connection()
    conn.execute("""
        INSERT INTO historique (materiel_id, code_immo, nom_materiel, action, details, utilisateur, date_action)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (materiel_id, code_immo, nom_materiel, action, details,
          session.get("username", "?"), maintenant))
    conn.commit()
    conn.close()


def envoyer_email(adresse, sujet, corps):
    """Envoie un email et renvoie (ok, message)."""
    if not SMTP_ACTIF:
        return False, "Email non configuré. Renseigne les variables SMTP_HOST, SMTP_USER, SMTP_PASS."
    if not adresse:
        return False, "Aucune adresse email enregistrée pour cette personne."
    try:
        msg = EmailMessage()
        msg["From"] = EMAIL_FROM
        msg["To"] = adresse
        msg["Subject"] = sujet
        msg.set_content(corps)
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as serveur:
            serveur.ehlo()
            if serveur.has_extn("starttls"):
                serveur.starttls()
                serveur.ehlo()
            serveur.login(SMTP_USER, SMTP_PASS)
            serveur.send_message(msg)
        return True, f"Email envoyé à {adresse}."
    except Exception as e:
        return False, f"Erreur lors de l'envoi de l'email : {e}"


def prets_en_cours():
    """Tous les prêts non rendus, enrichis du statut retard/aujourd'hui."""
    conn = get_db_connection()
    prets = conn.execute("SELECT * FROM prets WHERE statut != 'Rendu' ORDER BY date_fin").fetchall()
    conn.close()
    aujourd_hui = date.today()
    resultats = []
    for p in prets:
        fin = datetime.strptime(p["date_fin"], "%Y-%m-%d").date() if p["date_fin"] else None
        stats = {"en_retard": bool(fin and fin < aujourd_hui),
                 "aujourdhui": bool(fin and fin == aujourd_hui)}
        resultats.append((p, stats))
    return resultats


def alertes_stock():
    """Matériels dont il ne reste qu'1/6 du stock initial."""
    conn = get_db_connection()
    items = conn.execute("""
        SELECT * FROM materiel
        WHERE quantite > 0 AND quantite_initiale > 0
          AND quantite <= (quantite_initiale / 6)
        ORDER BY quantite ASC
    """).fetchall()
    conn.close()
    return items


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        if session.get("role") != "admin":
            flash("Accès réservé aux administrateurs.", "erreur")
            return redirect(url_for("accueil"))
        return f(*args, **kwargs)
    return decorated


# ---------- AUTHENTIFICATION ----------

@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("accueil"))

    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        conn = get_db_connection()
        user = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
        conn.close()

        if user and check_password_hash(user["password_hash"], password):
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user["role"]
            return redirect(url_for("accueil"))
        flash("Identifiant ou mot de passe incorrect.", "erreur")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Vous êtes déconnecté.", "succes")
    return redirect(url_for("login"))


# ---------- GESTION DES UTILISATEURS (ADMIN) ----------

@app.route("/utilisateurs")
@admin_required
def utilisateurs():
    conn = get_db_connection()
    users = conn.execute("SELECT * FROM users ORDER BY username").fetchall()
    conn.close()
    return render_template("utilisateurs.html", users=users)


@app.route("/utilisateurs/ajouter", methods=["POST"])
@admin_required
def utilisateurs_ajouter():
    username = request.form["username"].strip()
    password = request.form["password"]
    role = request.form.get("role", "user")

    conn = get_db_connection()
    try:
        conn.execute(
            "INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
            (username, generate_password_hash(password), role)
        )
        conn.commit()
        flash(f"Utilisateur '{username}' créé.", "succes")
    except DBIntegrityError:
        flash(f"L'identifiant '{username}' existe déjà.", "erreur")
    conn.close()
    return redirect(url_for("utilisateurs"))


@app.route("/utilisateurs/supprimer/<int:id>")
@admin_required
def utilisateurs_supprimer(id):
    if id == session.get("user_id"):
        flash("Tu ne peux pas supprimer ton propre compte.", "erreur")
        return redirect(url_for("utilisateurs"))
    conn = get_db_connection()
    conn.execute("DELETE FROM users WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    flash("Utilisateur supprimé.", "succes")
    return redirect(url_for("utilisateurs"))


# ---------- ACCUEIL ----------

@app.route("/")
@login_required
def accueil():
    conn = get_db_connection()
    nb_categories = conn.execute("SELECT COUNT(*) AS total FROM categories").fetchone()["total"]
    nb_materiel = conn.execute("SELECT COUNT(*) AS total FROM materiel").fetchone()["total"]

    repartition = conn.execute("""
        SELECT categories.nom AS nom, categories.icone AS icone, COUNT(materiel.id) AS nb
        FROM categories
        LEFT JOIN materiel ON materiel.categorie_id = categories.id
        GROUP BY categories.id
        ORDER BY nb DESC
    """).fetchall()
    max_nb = max([r["nb"] for r in repartition], default=1) or 1

    derniers = conn.execute("""
        SELECT materiel.*, categories.nom AS categorie_nom, categories.icone AS categorie_icone
        FROM materiel
        JOIN categories ON materiel.categorie_id = categories.id
        ORDER BY materiel.id DESC LIMIT 5
    """).fetchall()
    conn.close()

    prets_alertes = prets_en_cours()
    stocks = alertes_stock()

    return render_template(
        "accueil.html", nb_categories=nb_categories, nb_materiel=nb_materiel,
        repartition=repartition, max_nb=max_nb, derniers=derniers,
        prets_alertes=prets_alertes, stocks=stocks
    )


# ---------- SCANNER (code-barres / QR code) ----------

@app.route("/scanner")
@login_required
def scanner():
    conn = get_db_connection()
    categories = conn.execute("SELECT * FROM categories ORDER BY nom").fetchall()
    conn.close()
    return render_template("scanner.html", categories=categories)


@app.route("/scanner/verifier")
@login_required
def scanner_verifier():
    code = request.args.get("code", "").strip()
    categorie_id = request.args.get("categorie_id", type=int)

    if not code:
        flash("Aucun code reçu.", "erreur")
        return redirect(url_for("scanner"))

    conn = get_db_connection()
    existant = conn.execute(
        "SELECT id, categorie_id, nom FROM materiel WHERE code_immo = ?", (code,)
    ).fetchone()
    conn.close()

    if existant:
        flash(f"Le code '{code}' existe déjà ({existant['nom']}). Voici sa fiche.", "succes")
        return redirect(url_for("materiel_modifier", categorie_id=existant["categorie_id"], id=existant["id"]))

    if not categorie_id:
        flash("Choisis d'abord une catégorie pour enregistrer un nouveau matériel.", "erreur")
        return redirect(url_for("scanner"))

    flash(f"Code '{code}' inconnu : complète la fiche pour l'enregistrer.", "succes")
    return redirect(url_for("materiel_ajouter", categorie_id=categorie_id, code_immo=code))


# ---------- CATEGORIES ----------

@app.route("/materiel")
@login_required
def materiel_categories():
    conn = get_db_connection()
    categories = conn.execute("""
        SELECT categories.*, COUNT(materiel.id) AS nb_items
        FROM categories
        LEFT JOIN materiel ON materiel.categorie_id = categories.id
        GROUP BY categories.id
        ORDER BY categories.nom
    """).fetchall()
    conn.close()
    return render_template("materiel_categories.html", categories=categories)


@app.route("/materiel/ajouter_categorie", methods=["POST"])
@login_required
def ajouter_categorie():
    nom = request.form["nom"].strip()
    icone = request.form.get("icone", "📦")
    if nom:
        conn = get_db_connection()
        try:
            conn.execute("INSERT INTO categories (nom, icone) VALUES (?, ?)", (nom, icone))
            conn.commit()
            flash(f"Catégorie '{nom}' ajoutée.", "succes")
        except DBIntegrityError:
            flash(f"La catégorie '{nom}' existe déjà.", "erreur")
        conn.close()
    return redirect(url_for("materiel_categories"))


@app.route("/materiel/categorie/supprimer/<int:id>")
@admin_required
def supprimer_categorie(id):
    conn = get_db_connection()
    conn.execute("DELETE FROM categories WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    flash("Catégorie supprimée (et son matériel associé).", "succes")
    return redirect(url_for("materiel_categories"))


# ---------- MATERIEL D'UNE CATEGORIE ----------

@app.route("/materiel/<int:categorie_id>")
@login_required
def materiel_liste(categorie_id):
    conn = get_db_connection()
    categorie = conn.execute("SELECT * FROM categories WHERE id = ?", (categorie_id,)).fetchone()
    if categorie is None:
        conn.close()
        flash("Catégorie introuvable.", "erreur")
        return redirect(url_for("materiel_categories"))

    recherche = request.args.get("recherche", "")
    statut_filtre = request.args.get("statut", "")

    requete = "SELECT * FROM materiel WHERE categorie_id = ?"
    params = [categorie_id]

    if recherche:
        requete += " AND (nom LIKE ? OR code_immo LIKE ? OR type LIKE ? OR responsable LIKE ?)"
        params += [f"%{recherche}%", f"%{recherche}%", f"%{recherche}%", f"%{recherche}%"]

    if statut_filtre:
        requete += " AND statut = ?"
        params.append(statut_filtre)

    requete += " ORDER BY id DESC"
    items = conn.execute(requete, params).fetchall()
    conn.close()
    return render_template(
        "materiel_liste.html", categorie=categorie, items=items,
        recherche=recherche, statut_filtre=statut_filtre, statuts=STATUTS
    )


@app.route("/materiel/<int:categorie_id>/ajouter", methods=["GET", "POST"])
@login_required
def materiel_ajouter(categorie_id):
    conn = get_db_connection()
    categorie = conn.execute("SELECT * FROM categories WHERE id = ?", (categorie_id,)).fetchone()
    if categorie is None:
        conn.close()
        flash("Catégorie introuvable.", "erreur")
        return redirect(url_for("materiel_categories"))

    if request.method == "POST":
        type_ = request.form["type"]
        code_immo = request.form["code_immo"].strip()
        prix = request.form["prix"] or 0
        projet = request.form["projet"]
        nom = request.form["nom"]
        date_materiel = request.form["date_materiel"]
        statut = request.form.get("statut", "En service")
        responsable = request.form.get("responsable", "")
        fournisseur = request.form.get("fournisseur", "")
        quantite = int(request.form.get("quantite", 1) or 1)
        quantite_initiale = int(request.form.get("quantite_initiale", quantite) or quantite)

        try:
            nouveau_id = inserer(conn, """
                INSERT INTO materiel (categorie_id, type, code_immo, prix, projet, nom, date_materiel, statut, responsable, fournisseur, quantite, quantite_initiale)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (categorie_id, type_, code_immo, prix, projet, nom, date_materiel, statut, responsable, fournisseur, quantite, quantite_initiale))
            conn.commit()
            conn.close()
            log_historique(nouveau_id, code_immo, nom, "Ajout", f"Matériel ajouté dans '{categorie['nom']}'")
            flash("Matériel ajouté avec succès.", "succes")
            return redirect(url_for("materiel_liste", categorie_id=categorie_id))
        except DBIntegrityError:
            conn.close()
            flash(f"Le code IMMO '{code_immo}' existe déjà. Chaque matériel doit avoir un code IMMO unique.", "erreur")
            return render_template("materiel_ajouter.html", categorie=categorie, valeurs=request.form, statuts=STATUTS)

    conn.close()
    valeurs = {"code_immo": request.args.get("code_immo", "")}
    return render_template("materiel_ajouter.html", categorie=categorie, valeurs=valeurs, statuts=STATUTS)


@app.route("/materiel/<int:categorie_id>/modifier/<int:id>", methods=["GET", "POST"])
@login_required
def materiel_modifier(categorie_id, id):
    conn = get_db_connection()
    categorie = conn.execute("SELECT * FROM categories WHERE id = ?", (categorie_id,)).fetchone()

    if request.method == "POST":
        type_ = request.form["type"]
        code_immo = request.form["code_immo"].strip()
        prix = request.form["prix"] or 0
        projet = request.form["projet"]
        nom = request.form["nom"]
        date_materiel = request.form["date_materiel"]
        statut = request.form.get("statut", "En service")
        responsable = request.form.get("responsable", "")
        fournisseur = request.form.get("fournisseur", "")
        quantite = int(request.form.get("quantite", 1) or 1)
        quantite_initiale = int(request.form.get("quantite_initiale", quantite) or quantite)

        try:
            conn.execute("""
                UPDATE materiel
                SET type=?, code_immo=?, prix=?, projet=?, nom=?, date_materiel=?, statut=?, responsable=?, fournisseur=?, quantite=?, quantite_initiale=?
                WHERE id=?
            """, (type_, code_immo, prix, projet, nom, date_materiel, statut, responsable, fournisseur, quantite, quantite_initiale, id))
            conn.commit()
            conn.close()
            log_historique(id, code_immo, nom, "Modification", f"Statut: {statut}, Responsable: {responsable or '-'}")
            flash("Matériel mis à jour.", "succes")
            return redirect(url_for("materiel_liste", categorie_id=categorie_id))
        except DBIntegrityError:
            item = conn.execute("SELECT * FROM materiel WHERE id = ?", (id,)).fetchone()
            conn.close()
            flash(f"Le code IMMO '{code_immo}' est déjà utilisé par un autre matériel.", "erreur")
            return render_template("materiel_modifier.html", categorie=categorie, item=item, statuts=STATUTS)

    item = conn.execute("SELECT * FROM materiel WHERE id = ?", (id,)).fetchone()
    conn.close()
    return render_template("materiel_modifier.html", categorie=categorie, item=item, statuts=STATUTS)


@app.route("/materiel/<int:categorie_id>/supprimer/<int:id>")
@login_required
def materiel_supprimer(categorie_id, id):
    conn = get_db_connection()
    item = conn.execute("SELECT * FROM materiel WHERE id = ?", (id,)).fetchone()
    conn.execute("DELETE FROM materiel WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    if item:
        log_historique(None, item["code_immo"], item["nom"], "Suppression", "Matériel supprimé de l'inventaire")
    flash("Matériel supprimé.", "succes")
    return redirect(url_for("materiel_liste", categorie_id=categorie_id))


# ---------- HISTORIQUE ----------

@app.route("/historique")
@login_required
def historique():
    conn = get_db_connection()
    logs = conn.execute("SELECT * FROM historique ORDER BY id DESC LIMIT 300").fetchall()
    conn.close()
    return render_template("historique.html", logs=logs)


# ---------- GESTION DES PRETS ----------

@app.route("/prets")
@login_required
def prets():
    conn = get_db_connection()
    materiels = conn.execute("SELECT id, code_immo, nom FROM materiel ORDER BY nom").fetchall()
    prets = conn.execute("SELECT * FROM prets ORDER BY date_fin DESC, id DESC").fetchall()
    conn.close()

    prets_retard = []
    prets_aujourdhui = []
    prets_en_cours_liste = []
    prets_rendus = []
    aujourd_hui = date.today()
    for p in prets:
        fin = datetime.strptime(p["date_fin"], "%Y-%m-%d").date() if p["date_fin"] else None
        if p["statut"] == "Rendu":
            prets_rendus.append(p)
        else:
            prets_en_cours_liste.append(p)
            if fin and fin < aujourd_hui:
                prets_retard.append(p)
            if fin and fin == aujourd_hui:
                prets_aujourdhui.append(p)

    return render_template(
        "prets.html", materiels=materiels, prets_en_cours=prets_en_cours_liste,
        prets_rendus=prets_rendus, prets_retard=prets_retard,
        prets_aujourdhui=prets_aujourdhui
    )


@app.route("/prets/ajouter", methods=["POST"])
@login_required
def prets_ajouter():
    materiel_id = request.form["materiel_id"]
    emprunteur = request.form["emprunteur"].strip()
    email_emprunteur = request.form["email_emprunteur"].strip()
    date_debut = request.form["date_debut"]
    date_fin = request.form["date_fin"]

    if not emprunteur:
        flash("Indique le nom de l'emprunteur.", "erreur")
        return redirect(url_for("prets"))

    conn = get_db_connection()
    materiel = conn.execute("SELECT * FROM materiel WHERE id = ?", (materiel_id,)).fetchone()
    if materiel is None:
        conn.close()
        flash("Matériel introuvable.", "erreur")
        return redirect(url_for("prets"))

    inserer(conn, """
        INSERT INTO prets (materiel_id, code_immo, nom_materiel, emprunteur, email_emprunteur, date_debut, date_fin, statut)
        VALUES (?, ?, ?, ?, ?, ?, ?, 'En cours')
    """, (materiel_id, materiel["code_immo"], materiel["nom"], emprunteur, email_emprunteur, date_debut, date_fin))
    conn.commit()
    conn.close()
    log_historique(materiel["id"], materiel["code_immo"], materiel["nom"], "Prêt",
                   f"Prêté à {emprunteur} du {date_debut} au {date_fin}")
    flash(f"Prêt enregistré : {materiel['nom']} → {emprunteur} jusqu'au {date_fin}.", "succes")
    return redirect(url_for("prets"))


@app.route("/prets/rendre/<int:id>")
@login_required
def prets_rendre(id):
    conn = get_db_connection()
    pret = conn.execute("SELECT * FROM prets WHERE id = ?", (id,)).fetchone()
    if pret:
        aujourd_hui = date.today().isoformat()
        conn.execute("UPDATE prets SET statut = 'Rendu', date_retour = ? WHERE id = ?", (aujourd_hui, id))
        conn.commit()
        log_historique(pret["materiel_id"], pret["code_immo"], pret["nom_materiel"], "Retour",
                       f"Retour par {pret['emprunteur']} le {aujourd_hui}")
        flash(f"Retour enregistré pour {pret['nom_materiel']} ({pret['emprunteur']}).", "succes")
        en_retard = (pret["date_fin"] and pret["date_fin"] < aujourd_hui)
        if pret["email_emprunteur"] and en_retard:
            ok, m = envoyer_email(
                pret["email_emprunteur"],
                f"Retour enregistré - {pret['nom_materiel']}",
                f"Bonjour {pret['emprunteur']},\n\nLe matériel {pret['nom_materiel']} "
                f"(code {pret['code_immo']}) a bien été rendu le {aujourd_hui}.\nMerci."
            )
            flash(m, "succes" if ok else "erreur")
    conn.close()
    return redirect(url_for("prets"))


@app.route("/prets/email/<int:id>")
@login_required
def prets_email(id):
    conn = get_db_connection()
    pret = conn.execute("SELECT * FROM prets WHERE id = ?", (id,)).fetchone()
    conn.close()
    if not pret:
        flash("Prêt introuvable.", "erreur")
        return redirect(url_for("prets"))

    if pret["statut"] == "Rendu":
        sujet = f"Reçu pour la restitution - {pret['nom_materiel']}"
        corps = (f"Bonjour {pret['emprunteur']},\n\n"
                 f"Ceci est une confirmation : le matériel {pret['nom_materiel']} "
                 f"(code {pret['code_immo']}) a été rendu le {pret['date_retour']}.\n")
    else:
        sujet = f"Rappel : rendre {pret['nom_materiel']}"
        corps = (f"Bonjour {pret['emprunteur']},\n\n"
                 f"Tu as emprunté le matériel : {pret['nom_materiel']} (code {pret['code_immo']}).\n"
                 f"Date de début : {pret['date_debut']}\nDate de fin prévue : {pret['date_fin']}\n\n"
                 f"Si le délai est dépassé, merci de le rendre dès que possible.\nMerci.")
    ok, m = envoyer_email(pret["email_emprunteur"], sujet, corps)
    flash(m, "succes" if ok else "erreur")
    return redirect(url_for("prets"))


# ---------- EXPORTS PAR CATEGORIE ----------

def get_items_categorie(categorie_id):
    conn = get_db_connection()
    categorie = conn.execute("SELECT * FROM categories WHERE id = ?", (categorie_id,)).fetchone()
    items = conn.execute(
        "SELECT * FROM materiel WHERE categorie_id = ? ORDER BY id", (categorie_id,)
    ).fetchall()
    conn.close()
    return categorie, items


def get_tout_le_materiel():
    conn = get_db_connection()
    items = conn.execute("""
        SELECT materiel.*, categories.nom AS categorie_nom
        FROM materiel JOIN categories ON materiel.categorie_id = categories.id
        ORDER BY categories.nom, materiel.id
    """).fetchall()
    conn.close()
    return items


COLONNES_EXPORT = ["Catégorie", "Type", "Code IMMO", "Nom", "Prix", "Projet", "Statut", "Responsable", "Fournisseur", "Date"]


def ligne_export(m, avec_categorie=False):
    ligne = []
    if avec_categorie:
        ligne.append(m["categorie_nom"])
    ligne += [m["type"] or "", m["code_immo"], m["nom"], str(m["prix"] or ""), m["projet"] or "",
              m["statut"] or "", m["responsable"] or "", m["fournisseur"] or "", m["date_materiel"] or ""]
    return ligne


@app.route("/materiel/<int:categorie_id>/export/pdf")
@login_required
def export_pdf(categorie_id):
    categorie, items = get_items_categorie(categorie_id)
    return generer_pdf(f"Inventaire - {categorie['nom']}", COLONNES_EXPORT[1:], [ligne_export(m) for m in items], categorie["nom"])


@app.route("/materiel/<int:categorie_id>/export/word")
@login_required
def export_word(categorie_id):
    categorie, items = get_items_categorie(categorie_id)
    return generer_word(f"Inventaire - {categorie['nom']}", COLONNES_EXPORT[1:], [ligne_export(m) for m in items], categorie["nom"])


@app.route("/materiel/<int:categorie_id>/export/excel")
@login_required
def export_excel(categorie_id):
    categorie, items = get_items_categorie(categorie_id)
    return generer_excel(categorie["nom"], COLONNES_EXPORT[1:], [ligne_export(m) for m in items], categorie["nom"])


# ---------- EXPORT GLOBAL (toutes catégories) ----------

@app.route("/export/global/pdf")
@login_required
def export_global_pdf():
    items = get_tout_le_materiel()
    return generer_pdf("Inventaire complet - FID", COLONNES_EXPORT, [ligne_export(m, True) for m in items], "complet")


@app.route("/export/global/word")
@login_required
def export_global_word():
    items = get_tout_le_materiel()
    return generer_word("Inventaire complet - FID", COLONNES_EXPORT, [ligne_export(m, True) for m in items], "complet")


@app.route("/export/global/excel")
@login_required
def export_global_excel():
    items = get_tout_le_materiel()
    return generer_excel("Inventaire complet", COLONNES_EXPORT, [ligne_export(m, True) for m in items], "complet")


# ---------- FONCTIONS DE GENERATION DE FICHIERS ----------

def generer_pdf(titre, entetes, lignes, nom_fichier):
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements = [Paragraph(titre, styles["Title"]), Spacer(1, 12)]

    data = [entetes] + lignes
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123a6b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef3f8")]),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name=f"inventaire_{nom_fichier}.pdf", mimetype="application/pdf")


def generer_word(titre, entetes, lignes, nom_fichier):
    from docx import Document

    doc = Document()
    doc.add_heading(titre, level=1)

    table = doc.add_table(rows=1, cols=len(entetes))
    table.style = "Light Grid Accent 1"
    for i, t in enumerate(entetes):
        table.rows[0].cells[i].text = t

    for ligne in lignes:
        cells = table.add_row().cells
        for i, val in enumerate(ligne):
            cells[i].text = str(val)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name=f"inventaire_{nom_fichier}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def generer_excel(titre_feuille, entetes, lignes, nom_fichier):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = titre_feuille[:30]
    ws.append(entetes)
    for ligne in lignes:
        ws.append(ligne)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name=f"inventaire_{nom_fichier}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


init_db()

if __name__ == "__main__":
    app.run(debug=True)